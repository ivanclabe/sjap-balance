import json
import openpyxl
from openpyxl.utils import column_index_from_string as ci

wb = openpyxl.load_workbook('balance-ejemplo-feb-2026.xlsx', data_only=True)

def cell(ws, col, row):
    return ws.cell(row=row, column=ci(col)).value

def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(v, 4)
    return v

out = {}

# ---- Producto sheets: G. Corriente / Bioacem / G. Extra ----
PRODUCT_SHEETS = {'G. Corriente': 'CORRIENTE', 'Bioacem': 'BIOACEM', 'G. Extra': 'EXTRA'}
producto_dias = {}
for sheet_name, codigo in PRODUCT_SHEETS.items():
    ws = wb[sheet_name]
    dias = []
    for r in range(7, 38):
        dia = cell(ws, 'B', r)
        if dia is None or dia > 28:
            continue
        dias.append({
            'dia': dia,
            'inventario_inicial': num(cell(ws, 'D', r)),
            'ventas': num(cell(ws, 'E', r)),
            'recibos': num(cell(ws, 'F', r)),
            'inventario_teorico': num(cell(ws, 'G', r)),
            'inventario_final': num(cell(ws, 'H', r)),
            'fluctuacion_dia': num(cell(ws, 'I', r)),
            'fluctuacion_acum': num(cell(ws, 'J', r)),
            'fluctuacion_valor': num(cell(ws, 'K', r)),
        })
    producto_dias[codigo] = dias
out['producto_dias'] = producto_dias

# ---- Cierre ----
ws = wb['Cierre']
cierre_cols = {
    'dia': 'B', 'vta_corriente': 'D', 'vta_bioacem': 'E', 'vta_extra': 'F', 'vta_total': 'G',
    'vta_gl': 'H', 'num_clientes': 'I', 'efectivo': 'J', 'rumbo': 'L', 'urea_rumbo': 'N',
    'cons_terpel': 'P', 'mi_emp': 'R', 'clien_propios': 'T', 'app_terpel': 'V', 'go_pass': 'X',
    'datafono': 'Z', 'qr': 'AB', 'ajustes': 'AD', 'vive_terpel': 'AF', 'urea': 'AH',
    'efectivo_real': 'AJ', 'diferencia_efectivo': 'AK', 'diferencia_efectivo_acum': 'AL',
    'certificacion': 'AM', 'diferencia_certificacion': 'AN', 'diferencia_certificacion_acum': 'AO',
}
cierre_dias = []
for r in range(3, 34):
    dia = cell(ws, 'B', r)
    if dia is None or dia > 28:
        continue
    row = {k: num(cell(ws, col, r)) for k, col in cierre_cols.items()}
    cierre_dias.append(row)
out['cierre_dias'] = cierre_dias

# ---- Vtas efectivo (cuentas de clientes propios) ----
ws = wb['Vtas efectivo']
account_cols = []
for c in range(ci('D'), ci('BI') + 1):
    name = ws.cell(row=2, column=c).value
    if name:
        account_cols.append((c, name))
vtas_dias = []
for r in range(3, 34):
    dia = cell(ws, 'B', r)
    if dia is None or dia > 28:
        continue
    movimientos = []
    for c, name in account_cols:
        v = ws.cell(row=r, column=c).value
        if v:
            movimientos.append({'cuenta': name, 'monto': num(v)})
    vtas_dias.append({
        'dia': dia,
        'total_combustible': num(cell(ws, 'BJ', r)),
        'total_urea': num(cell(ws, 'BK', r)),
        'total_lub': num(cell(ws, 'BL', r)),
        'total': num(cell(ws, 'BM', r)),
        'movimientos': movimientos,
    })
out['cuentas_nombres'] = [n for _, n in account_cols]
out['vtas_efectivo_dias'] = vtas_dias

# ---- Urea ----
ws = wb['Urea']
urea_cols = {
    'dia': 'B', 'inventario_inicial': 'D', 'inventario_final': 'E', 'diferencia': 'F',
    'rumbo_litros': 'G', 'clientes_propios_litros': 'H', 'valor': 'I', 'valor_rumbo': 'J',
    'valor_propios': 'K', 'total': 'L', 'inventario_teorico': 'M', 'recibo': 'N',
}
urea_dias = []
for r in range(4, 32):
    dia = cell(ws, 'B', r)
    if dia is None or dia > 28:
        continue
    row = {k: num(cell(ws, col, r)) for k, col in urea_cols.items()}
    urea_dias.append(row)
out['urea_dias'] = urea_dias
out['urea_costo_unitario'] = num(cell(ws, 'R', 3))

# ---- Facturas ----
ws = wb['Facturas']
facturas = []
for r in range(4, 60):
    fecha = cell(ws, 'B', r)
    factura = cell(ws, 'C', r)
    cte = cell(ws, 'D', r)
    extra = cell(ws, 'E', r)
    bioacem = cell(ws, 'G', r)
    if fecha is None:
        continue
    if not (factura or cte or extra or bioacem):
        continue
    facturas.append({
        'fecha': str(fecha)[:10],
        'numero_factura': factura,
        'corriente': num(cte),
        'extra': num(extra),
        'bioacem': num(bioacem),
    })
out['facturas'] = facturas

# ---- Presupuesto (columna FEB = D) ----
ws = wb['Presupuesto']
presupuesto_rows = {}
for r in range(7, 20):
    label = cell(ws, 'B', r)
    if not label:
        continue
    presupuesto_rows[label] = num(cell(ws, 'D', r))
out['presupuesto_feb'] = presupuesto_rows

# ---- Sicom ----
ws = wb['Sicom']
sicom = []
for r in range(9, 12):
    producto = cell(ws, 'B', r)
    sicom.append({
        'producto': producto,
        'inventario_inicial': num(cell(ws, 'C', r)),
        'compras': num(cell(ws, 'D', r)),
        'ventas': num(cell(ws, 'E', r)),
        'faltantes': num(cell(ws, 'F', r)),
        'evaporacion': num(cell(ws, 'G', r)),
        'inventario_final_calc': num(cell(ws, 'H', r)),
        'inventario_final_real': num(cell(ws, 'I', r)),
    })
out['sicom'] = sicom

# ---- Promotor (detalle turno/isla) ----
ws = wb['Promotor']
# fila 6: nombres de promotor (uno cada bloque de 6 columnas), fila 7: subencabezados
promotor_blocks = []
c = ci('D')
while c <= ci('CM'):
    nombre = ws.cell(row=6, column=c).value
    if nombre:
        promotor_blocks.append((nombre, c))
    c += 6
promotor_dias = []
for r in range(8, 39):
    dia = cell(ws, 'B', r)
    if dia is None or dia > 28:
        continue
    weekday = cell(ws, 'C', r)
    turnos = []
    for nombre, startcol in promotor_blocks:
        vals = {label: ws.cell(row=r, column=startcol + i).value
                for i, label in enumerate(['T1', 'T2', 'T3', 'T4', 'NoVEH', 'ISLA'])}
        num_veh = vals['NoVEH']
        isla = vals['ISLA']
        for label in ('T1', 'T2', 'T3', 'T4'):
            v = vals[label]
            if v is not None:
                turnos.append({
                    'promotor': nombre, 'turno': label, 'galones': num(v),
                    'num_vehiculos': num(num_veh), 'isla': str(isla) if isla is not None else None,
                })
    promotor_dias.append({'dia': dia, 'weekday': weekday, 'turnos': turnos})
out['promotor_dias'] = promotor_dias

# Totales de promotor (filas 45-57 aprox)
promotor_totales = []
for r in range(45, 58):
    nombre = cell(ws, 'D', r)
    if not nombre:
        continue
    promotor_totales.append({
        'promotor': nombre,
        'galones': num(cell(ws, 'F', r)),
        'clientes': num(cell(ws, 'H', r)),
    })
out['promotor_totales'] = promotor_totales

with open('extracted.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=1, default=str)

print('OK')
print('producto_dias:', {k: len(v) for k, v in producto_dias.items()})
print('cierre_dias:', len(cierre_dias))
print('cuentas:', len(out['cuentas_nombres']))
print('vtas_efectivo_dias:', len(vtas_dias))
print('urea_dias:', len(urea_dias))
print('facturas:', len(facturas))
print('presupuesto_feb keys:', list(presupuesto_rows.keys()))
print('sicom:', len(sicom))
print('promotor_blocks:', len(promotor_blocks), [n for n,_ in promotor_blocks])
print('promotor_dias:', len(promotor_dias))
print('promotor_totales:', len(promotor_totales))
